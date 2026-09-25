"""
Structured XLSX Report Generator.
Utilizes openpyxl to generate multi-sheet spreadsheet audits containing
all findings, summary metrics, conditional formatting colors, and rule pack references.
"""

from datetime import datetime, timezone
import io
from typing import Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..ai.schemas import QCAnalysisResult, SeverityEnum


class XLSXReportGenerator:
    """Compiles QC findings into formatted Excel workbooks for ERP/PLM integration."""

    HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    REGULAR_FONT = Font(name="Calibri", size=10)
    BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Severity-specific color fills for high-impact visual categorization
    SEVERITY_FILLS = {
        SeverityEnum.CRITICAL: PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        SeverityEnum.MAJOR: PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        SeverityEnum.MINOR: PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        SeverityEnum.INFO: PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
    }
    SEVERITY_FONTS = {
        SeverityEnum.CRITICAL: Font(name="Calibri", size=10, bold=True, color="991B1B"),
        SeverityEnum.MAJOR: Font(name="Calibri", size=10, bold=True, color="9A3412"),
        SeverityEnum.MINOR: Font(name="Calibri", size=10, bold=True, color="854D0E"),
        SeverityEnum.INFO: Font(name="Calibri", size=10, bold=True, color="075985"),
    }

    def generate(self, result: QCAnalysisResult, output_destination: Union[str, io.BytesIO]):
        wb = openpyxl.Workbook()

        # Sheet 1: Executive Summary
        ws_sum = wb.active
        ws_sum.title = "QC Summary"
        self._build_summary_sheet(ws_sum, result)

        # Sheet 2: Discrepancy Findings Matrix
        ws_findings = wb.create_sheet(title="Discrepancy Details")
        self._build_findings_sheet(ws_findings, result)

        # Sheet 3: Standard Rules Reference
        ws_rules = wb.create_sheet(title="Standards & Rules Reference")
        self._build_rules_sheet(ws_rules, result)

        wb.save(output_destination)

    def _build_summary_sheet(self, ws, result: QCAnalysisResult):
        ws.append(["WIRING DIAGRAM QC ASSISTANT — AUDIT SUMMARY"])
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        ws.append(["Spandsons Horizon Engineering Pvt. Ltd. | Automated Compliance Inspection"])
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")
        ws.append([])

        metadata_rows = [
            ("Document ID", result.document_id),
            ("Filename", result.filename),
            ("Overall Compliance Status", result.overall_status.value),
            ("Applied Standards", ", ".join(result.standards_applied)),
            ("Ruleset Version", result.rules_version),
            ("Model / Engine Version", result.model_version),
            ("Processing Duration (ms)", result.processing_time_ms),
            ("Audit Timestamp (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ]
        for label, val in metadata_rows:
            ws.append([label, val])
            row_idx = ws.max_row
            ws.cell(row=row_idx, column=1).font = self.BOLD_FONT
            ws.cell(row=row_idx, column=2).font = self.REGULAR_FONT

        ws.append([])
        ws.append(["Audit Metric", "Count"])
        header_row = ws.max_row
        for col in range(1, 3):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT

        summary_metrics = [
            ("Total Checks Performed", result.summary.checks_total),
            ("Passed Checks", result.summary.passed),
            ("Failed Checks", result.summary.failed),
            ("Review Required", result.summary.review),
            ("Critical Findings", result.summary.critical_count),
            ("Major Findings", result.summary.major_count),
            ("Minor Findings", result.summary.minor_count),
            ("Info Findings", result.summary.info_count),
        ]
        for metric, count in summary_metrics:
            ws.append([metric, count])
            r = ws.max_row
            ws.cell(row=r, column=1).border = self.BORDER
            ws.cell(row=r, column=2).border = self.BORDER

        self._autofit_columns(ws)

    def _build_findings_sheet(self, ws, result: QCAnalysisResult):
        headers = [
            "Finding ID",
            "Page",
            "Severity",
            "Category",
            "Confidence Level",
            "Confidence Score",
            "Discrepancy Description",
            "Observed Evidence",
            "Engineering Requirement",
            "Standard Code",
            "Clause / Section",
            "Actionable Recommendation",
        ]
        ws.append(headers)

        header_row = 1
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for f in result.findings:
            row_data = [
                f.id,
                f.page,
                f.severity.value,
                f.category,
                f.confidence.level.value,
                round(f.confidence.score, 3),
                f.description,
                f.evidence,
                f.requirement,
                f.standard,
                f.standard_section,
                f.recommendation,
            ]
            ws.append(row_data)
            r = ws.max_row

            # Apply cell borders and severity color highlighting
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = self.BORDER
                cell.font = self.REGULAR_FONT

                # Color-code Severity Column (Column 3)
                if c == 3:
                    cell.fill = self.SEVERITY_FILLS.get(f.severity, PatternFill(fill_type=None))
                    cell.font = self.SEVERITY_FONTS.get(f.severity, self.BOLD_FONT)
                    cell.alignment = Alignment(horizontal="center")
                elif c in [1, 2, 5]:
                    cell.alignment = Alignment(horizontal="center")

        self._autofit_columns(ws)

    def _build_rules_sheet(self, ws, result: QCAnalysisResult):
        ws.append(["Standard Ruleset & Compliance Reference"])
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        ws.append([])

        headers = ["Rule ID", "Standard", "Section", "Category", "Mandatory Engineering Requirement"]
        ws.append(headers)
        h_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=col_idx)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT

        # Collect distinct rules invoked in the findings
        seen_rules = set()
        for f in result.findings:
            if f.rule_id not in seen_rules:
                seen_rules.add(f.rule_id)
                ws.append([f.rule_id, f.standard, f.standard_section, f.category, f.requirement])
                r = ws.max_row
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).border = self.BORDER

        self._autofit_columns(ws)

    def _autofit_columns(self, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(50, max(12, max_len + 3))
