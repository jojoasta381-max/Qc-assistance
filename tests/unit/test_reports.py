"""
Unit and API integration tests for Phase 8 Reporting.
Tests:
- build_analysis_result_from_db mapping logic
- PDFReportGenerator audit-grade layout, branding, and sign-off block
- XLSXReportGenerator 3-sheet structure, conditional formatting, and autofit
- /api/v1/qc-runs/{id}/report/summary JSON endpoint
- /api/v1/qc-runs/{id}/report/pdf Streaming PDF download
- /api/v1/qc-runs/{id}/report/xlsx Streaming Excel download
- /api/v1/qc-runs/{id}/findings/{finding_id} Single finding detail endpoint
- Multi-tenant IDOR isolation on report endpoints
"""

import io
import os
import tempfile
import uuid
import openpyxl
from pypdf import PdfReader
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from backend.src.ai.schemas import (
    BoundingBox,
    Confidence,
    ConfidenceLevelEnum,
    OverallStatusEnum,
    QCAnalysisResult,
    QCFinding as SchemaQCFinding,
    QCSummary,
    SeverityEnum,
)
from backend.src.api.deps import get_db, get_queue
from backend.src.api.main import app
from backend.src.infrastructure.database import Base
from backend.src.infrastructure.models import Document, Organization, QCFinding, QCRun, User
from backend.src.infrastructure.queue import AsyncInMemoryQueue
from backend.src.reports.builder import build_analysis_result_from_db
from backend.src.reports.pdf_generator import PDFReportGenerator
from backend.src.reports.xlsx_generator import XLSXReportGenerator
from backend.src.services.qc_worker import process_qc_job
from tests.unit.test_engine import create_sample_schematic_pdf


def test_build_analysis_result_from_db():
    """Verify ORM to domain model transformation logic."""
    run = QCRun(
        id=str(uuid.uuid4()),
        organization_id=str(uuid.uuid4()),
        document_id=str(uuid.uuid4()),
        overall_status="FAIL",
        checks_total=15,
        checks_passed=11,
        checks_failed=3,
        checks_review=1,
        model_version="qc-test-v1",
        rules_version="rules-test-v1",
        processing_time_ms=250,
    )
    doc = Document(
        id=run.document_id,
        organization_id=run.organization_id,
        project_id=str(uuid.uuid4()),
        filename="flight_control_schematic_rev_c.pdf",
        storage_path="/tmp/test.pdf",
        sha256_checksum="abc",
        file_size_bytes=1024,
        mime_type="application/pdf",
        page_count=1,
    )
    f1 = QCFinding(
        id=str(uuid.uuid4()),
        qc_run_id=run.id,
        finding_code="D-001",
        rule_id="RULE-WG-001",
        category="wire",
        description="Wire W101 missing gauge callout",
        severity="CRITICAL",
        confidence_level="HIGH",
        confidence_score=0.98,
        page_number=1,
        location_bbox={"x": 100, "y": 200, "width": 80, "height": 30},
        evidence_text="W101 (J1 to J2)",
        requirement_text="Wire gauge must be explicitly specified per IPC-WHMA-A-620D",
        standard_citation="IPC-WHMA-A-620D Section 4.1",
        recommendation="Add AWG callout to wire label.",
    )
    f2 = QCFinding(
        id=str(uuid.uuid4()),
        qc_run_id=run.id,
        finding_code="D-002",
        rule_id="RULE-CC-003",
        category="wire",
        description="Ambiguous conductor color code",
        severity="MAJOR",
        confidence_level="MEDIUM",
        confidence_score=0.85,
        page_number=2,
        location_bbox=None,
        evidence_text="W102 BL/WH",
        requirement_text="Standard 3-letter abbreviation required",
        standard_citation="UL 508A Section 66.5",
        recommendation="Replace with standard BLK or WHT abbreviation.",
    )

    result = build_analysis_result_from_db(run, [f1, f2], doc)

    assert result.document_id == run.document_id
    assert result.filename == "flight_control_schematic_rev_c.pdf"
    assert result.overall_status == OverallStatusEnum.FAIL
    assert len(result.findings) == 2
    assert result.findings[0].id == "D-001"
    assert result.findings[0].severity == SeverityEnum.CRITICAL
    assert result.findings[0].location.x == 100
    assert result.findings[1].id == "D-002"
    assert result.findings[1].severity == SeverityEnum.MAJOR
    assert result.summary.critical_count == 1
    assert result.summary.major_count == 1


def test_pdf_report_generator_formatting_and_sections():
    """Verify PDF layout includes header, summary, findings catalog, and sign-off."""
    summary = QCSummary(
        checks_total=12,
        passed=9,
        failed=2,
        review=1,
        critical_count=1,
        major_count=1,
        minor_count=0,
        info_count=0,
    )
    finding = SchemaQCFinding(
        id="D-001",
        rule_id="RULE-WG-001",
        category="wire",
        description="Wire W1 lacks AWG gauge definition.",
        severity=SeverityEnum.CRITICAL,
        confidence=Confidence(level=ConfidenceLevelEnum.HIGH, score=0.99),
        page=1,
        location=BoundingBox(x=50, y=50, width=100, height=40),
        evidence="W1 BLK",
        requirement="Gauge declaration mandatory.",
        standard="IPC-WHMA-A-620D",
        standard_section="Section 4.1",
        recommendation="Specify AWG 20.",
    )
    result = QCAnalysisResult(
        document_id="doc-12345",
        filename="avionics_wiring_harness.pdf",
        standards_applied=["IPC-WHMA-A-620D", "UL 508A"],
        overall_status=OverallStatusEnum.FAIL,
        summary=summary,
        findings=[finding],
        model_version="qc-hybrid-engine-v1.0",
        prompt_version="prompt-v1.0",
        rules_version="ruleset-v1.0",
        processing_time_ms=180,
    )

    buf = io.BytesIO()
    pdf_gen = PDFReportGenerator()
    pdf_gen.generate(result, buf)
    buf.seek(0)

    reader = PdfReader(buf)
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()

    # Spandsons Horizon Engineering Branding
    assert "SPANDSONS HORIZON ENGINEERING" in text
    assert "WIRING DIAGRAM QC ASSISTANT" in text
    assert "STATUS: FAIL" in text

    # Executive Summary & Metrics
    assert "Quality Control Executive Summary" in text
    assert "Total Checks" in text

    # Findings Catalog
    assert "Discrepancy Findings Catalog" in text
    assert "D-001" in text
    assert "CRITICAL" in text
    assert "Wire W1 lacks AWG gauge definition" in text

    # Engineering Sign-Off
    assert "3. Engineering Sign-Off & Verification" in text
    assert "Prepared By:" in text
    assert "Reviewed By:" in text


def test_pdf_report_generator_clean_pass():
    """Verify PDF rendering when 0 discrepancies are detected."""
    summary = QCSummary(
        checks_total=10,
        passed=10,
        failed=0,
        review=0,
        critical_count=0,
        major_count=0,
        minor_count=0,
        info_count=0,
    )
    result = QCAnalysisResult(
        document_id="doc-perfect",
        filename="perfect_harness.pdf",
        standards_applied=["IPC-WHMA-A-620D"],
        overall_status=OverallStatusEnum.PASS,
        summary=summary,
        findings=[],
        model_version="qc-hybrid-engine-v1.0",
        prompt_version="prompt-v1.0",
        rules_version="ruleset-v1.0",
        processing_time_ms=90,
    )

    buf = io.BytesIO()
    pdf_gen = PDFReportGenerator()
    pdf_gen.generate(result, buf)
    buf.seek(0)

    reader = PdfReader(buf)
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "STATUS: PASS" in text
    assert "No discrepancies or compliance violations detected" in text


def test_xlsx_report_generator_structure_and_sheets():
    """Verify multi-sheet XLSX output, conditional severity colors, and standards rules reference."""
    summary = QCSummary(
        checks_total=14,
        passed=11,
        failed=2,
        review=1,
        critical_count=1,
        major_count=1,
        minor_count=0,
        info_count=0,
    )
    f1 = SchemaQCFinding(
        id="D-001",
        rule_id="RULE-WG-001",
        category="wire",
        description="Wire missing gauge specification",
        severity=SeverityEnum.CRITICAL,
        confidence=Confidence(level=ConfidenceLevelEnum.HIGH, score=0.99),
        page=1,
        location=None,
        evidence="W100 RED",
        requirement="Wire gauge must be provided",
        standard="IPC-WHMA-A-620D",
        standard_section="Section 4.1",
        recommendation="Add 18 AWG callout",
    )
    f2 = SchemaQCFinding(
        id="D-002",
        rule_id="RULE-CC-003",
        category="wire",
        description="Color code ambiguity",
        severity=SeverityEnum.MAJOR,
        confidence=Confidence(level=ConfidenceLevelEnum.MEDIUM, score=0.88),
        page=1,
        location=None,
        evidence="W101 BL/WH",
        requirement="Single clear insulation color",
        standard="UL 508A",
        standard_section="Section 66.5",
        recommendation="Use BLK or WHT",
    )
    result = QCAnalysisResult(
        document_id="doc-xlsx-test",
        filename="chassis_wiring.pdf",
        standards_applied=["IPC-WHMA-A-620D", "UL 508A"],
        overall_status=OverallStatusEnum.FAIL,
        summary=summary,
        findings=[f1, f2],
        model_version="qc-v1",
        prompt_version="prompt-v1",
        rules_version="rules-v1",
        processing_time_ms=130,
    )

    buf = io.BytesIO()
    xlsx_gen = XLSXReportGenerator()
    xlsx_gen.generate(result, buf)
    buf.seek(0)

    wb = openpyxl.load_workbook(buf)
    expected_sheets = ["QC Summary", "Discrepancy Details", "Standards & Rules Reference"]
    assert all(sheet in wb.sheetnames for sheet in expected_sheets)

    # Sheet 1 checks
    ws_sum = wb["QC Summary"]
    assert ws_sum["A1"].value == "WIRING DIAGRAM QC ASSISTANT — AUDIT SUMMARY"
    assert ws_sum["A2"].value == "Spandsons Horizon Engineering Pvt. Ltd. | Automated Compliance Inspection"

    # Sheet 2 checks
    ws_find = wb["Discrepancy Details"]
    assert ws_find.max_row == 3  # Header + 2 findings
    assert ws_find.cell(row=2, column=1).value == "D-001"
    assert ws_find.cell(row=2, column=3).value == "CRITICAL"
    # Verify critical severity cell has color fill (FEE2E2)
    sev_cell = ws_find.cell(row=2, column=3)
    assert sev_cell.fill.start_color.rgb in ["00FEE2E2", "FEE2E2"]

    # Sheet 3 checks: Title, Blank, Header, Rule 1, Rule 2 -> max_row = 5
    ws_rules = wb["Standards & Rules Reference"]
    assert ws_rules.max_row == 5
    assert ws_rules.cell(row=4, column=1).value == "RULE-WG-001"
    assert ws_rules.cell(row=5, column=1).value == "RULE-CC-003"


# --- API Integration Fixtures & Tests ---

@pytest.fixture
async def qc_reports_api_env():
    """Setup in-memory DB and authenticated client with a completed QC run."""
    engine = create_async_engine(
        "sqlite+aiosqlite:///:memory:",
        connect_args={"check_same_thread": False},
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(bind=engine, class_=AsyncSession, expire_on_commit=False)
    test_queue = AsyncInMemoryQueue()

    async def override_get_db():
        async with session_factory() as session:
            yield session

    def override_get_queue():
        return test_queue

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_queue] = override_get_queue

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Register Org A
        reg = await client.post(
            "/api/v1/auth/register",
            json={
                "organization_name": "Spandsons Horizon Engineering",
                "organization_slug": "spandsons-eng",
                "full_name": "Pravin Chief Engineer",
                "email": "pravin@spandsons.com",
                "password": "Password123!",
            },
        )
        token_a = reg.json()["access_token"]
        client.headers.update({"Authorization": f"Bearer {token_a}"})

        # 2. Create Project
        proj = await client.post("/api/v1/projects", json={"name": "Harness Automation Rev 1"})
        project_id = proj.json()["id"]

        # 3. Create Sample Schematic PDF and Confirm Document
        with tempfile.TemporaryDirectory() as tmpdir:
            sample_pdf = os.path.join(tmpdir, "industrial_harness.pdf")
            create_sample_schematic_pdf(sample_pdf)
            doc_id = "doc_report_test_001"

            doc_resp = await client.post(
                f"/api/v1/documents/{doc_id}/confirm",
                json={
                    "project_id": project_id,
                    "filename": "industrial_harness.pdf",
                    "storage_path": sample_pdf,
                    "file_size_bytes": os.path.getsize(sample_pdf),
                    "mime_type": "application/pdf",
                    "sha256_checksum": "checksum9876543210fedcba",
                    "page_count": 1,
                },
            )
            assert doc_resp.status_code == 201

            # 4. Trigger QC Run
            run_res = await client.post(
                "/api/v1/qc-runs",
                json={
                    "document_id": doc_id,
                    "standards": ["IPC-WHMA-A-620D", "UL 508A"],
                    "enable_ai": True,
                },
            )
            assert run_res.status_code == 201
            run_id = run_res.json()["id"]

            # 5. Process QC job through worker
            job = await test_queue.dequeue()
            assert job is not None
            await process_qc_job(job, session_factory)

            # 6. Register Org B for cross-tenant IDOR verification
            reg_b = await client.post(
                "/api/v1/auth/register",
                json={
                    "organization_name": "Competitor Aerospace",
                    "organization_slug": "competitor-aero",
                    "full_name": "Intruder User",
                    "email": "intruder@competitor.com",
                    "password": "Password123!",
                },
            )
            token_b = reg_b.json()["access_token"]

            yield {
                "client": client,
                "token_a": token_a,
                "token_b": token_b,
                "run_id": run_id,
                "doc_id": doc_id,
            }

    app.dependency_overrides.clear()
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
    await engine.dispose()


@pytest.mark.anyio
async def test_qc_report_summary_endpoint(qc_reports_api_env):
    """Test GET /api/v1/qc-runs/{id}/report/summary returns structured executive metrics."""
    env = qc_reports_api_env
    client = env["client"]
    run_id = env["run_id"]
    client.headers.update({"Authorization": f"Bearer {env['token_a']}"})

    res = await client.get(f"/api/v1/qc-runs/{run_id}/report/summary")
    assert res.status_code == 200
    data = res.json()

    assert data["qc_run_id"] == run_id
    assert data["filename"] == "industrial_harness.pdf"
    assert data["overall_status"] in ["FAIL", "REVIEW_REQUIRED"]
    assert "standards_applied" in data
    assert data["total_findings"] >= 2
    assert "CRITICAL" in data["severity_breakdown"]
    assert "MAJOR" in data["severity_breakdown"]
    assert data["checks_summary"]["total"] >= 10


@pytest.mark.anyio
async def test_qc_report_pdf_download_endpoint(qc_reports_api_env):
    """Test GET /api/v1/qc-runs/{id}/report/pdf streams valid PDF artifact."""
    env = qc_reports_api_env
    client = env["client"]
    run_id = env["run_id"]
    client.headers.update({"Authorization": f"Bearer {env['token_a']}"})

    res = await client.get(f"/api/v1/qc-runs/{run_id}/report/pdf")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.pdf"')

    # Verify PDF content integrity
    pdf_bytes = io.BytesIO(res.content)
    reader = PdfReader(pdf_bytes)
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "SPANDSONS HORIZON ENGINEERING" in text
    assert "WIRING DIAGRAM QC ASSISTANT" in text


@pytest.mark.anyio
async def test_qc_report_xlsx_download_endpoint(qc_reports_api_env):
    """Test GET /api/v1/qc-runs/{id}/report/xlsx streams valid Excel matrix."""
    env = qc_reports_api_env
    client = env["client"]
    run_id = env["run_id"]
    client.headers.update({"Authorization": f"Bearer {env['token_a']}"})

    res = await client.get(f"/api/v1/qc-runs/{run_id}/report/xlsx")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.xlsx"')

    # Verify Excel workbook integrity
    xlsx_bytes = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(xlsx_bytes)
    assert "QC Summary" in wb.sheetnames
    assert "Discrepancy Details" in wb.sheetnames
    assert "Standards & Rules Reference" in wb.sheetnames


@pytest.mark.anyio
async def test_qc_finding_detail_endpoint(qc_reports_api_env):
    """Test GET /api/v1/qc-runs/{id}/findings/{finding_id}."""
    env = qc_reports_api_env
    client = env["client"]
    run_id = env["run_id"]
    client.headers.update({"Authorization": f"Bearer {env['token_a']}"})

    # Fetch findings list first
    list_res = await client.get(f"/api/v1/qc-runs/{run_id}/findings")
    assert list_res.status_code == 200
    findings = list_res.json()
    assert len(findings) > 0

    finding = findings[0]
    finding_id = finding["id"]
    finding_code = finding["finding_code"]

    # Query by UUID
    res_by_id = await client.get(f"/api/v1/qc-runs/{run_id}/findings/{finding_id}")
    assert res_by_id.status_code == 200
    data = res_by_id.json()
    assert data["id"] == finding_id
    assert data["finding_code"] == finding_code
    assert "evidence_text" in data
    assert "requirement_text" in data

    # Query by finding_code (e.g. D-001)
    res_by_code = await client.get(f"/api/v1/qc-runs/{run_id}/findings/{finding_code}")
    assert res_by_code.status_code == 200
    assert res_by_code.json()["id"] == finding_id


@pytest.mark.anyio
async def test_cross_tenant_idor_report_protection(qc_reports_api_env):
    """Verify Org B cannot download reports or access findings from Org A."""
    env = qc_reports_api_env
    client = env["client"]
    run_id = env["run_id"]

    # Switch to Org B
    client.headers.update({"Authorization": f"Bearer {env['token_b']}"})

    # 1. Summary
    res_summary = await client.get(f"/api/v1/qc-runs/{run_id}/report/summary")
    assert res_summary.status_code == 404

    # 2. PDF Download
    res_pdf = await client.get(f"/api/v1/qc-runs/{run_id}/report/pdf")
    assert res_pdf.status_code == 404

    # 3. Excel Download
    res_xlsx = await client.get(f"/api/v1/qc-runs/{run_id}/report/xlsx")
    assert res_xlsx.status_code == 404

    # 4. Finding Detail
    res_finding = await client.get(f"/api/v1/qc-runs/{run_id}/findings/D-001")
    assert res_finding.status_code == 404
